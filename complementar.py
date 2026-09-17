# -*- coding: utf-8 -*-
"""
Layout complementar: o layout gerado, em Excel, para o cliente completar.

Fluxo: ler XMLs -> aplicar De/Para -> gerar o complementar -> cliente preenche
-> importar no Senior. E o proprio resultado do writer, com o De/Para ja
aplicado, e nao uma lista separada de campos: o cliente ve a linha inteira e
completa so o que e dele.

Quem preenche cada campo vem da coluna 'Quem preenche' da planilha de
parametro. Enquanto ela nao existir, o papel e DERIVADO da acao do campo --
e esta derivacao e provisoria, nao decisao:

  eSocial  -> o extrator trouxe (DIRETO, DEPARA_*, REGRA, SEQUENCIAL)
  Fixo     -> valor fixo ou enviar vazio (FIXO, VAZIO)
  Cliente  -> sem origem no XML (SEM_ORIGEM, RELATORIO, PENDENTE)

Alem da coluna do cliente, a celula de um campo obrigatorio que o extrator
deveria ter trazido e saiu vazia tambem e destacada: o dado faltou naquele XML
e alguem precisa completar.

Layout 100% complementar de colaborador (nenhum campo vem do eSocial e ha
NUMCAD: 1016, 1017, 1022, 1024, 1025, 1035) sai com uma linha por colaborador do
1012, com NUMEMP, TIPCOL e NUMCAD ja preenchidos -- os mesmos valores do
cadastro, para o cliente so completar o resto. Cadastros de tabela (bancos,
locais, motivos) nao tem colaborador e continuam em branco.
"""
import io

PAPEL_POR_ACAO = {
    "DIRETO": "eSocial", "DEPARA_GERAL": "eSocial", "DEPARA_CLIENTE": "eSocial",
    "DEPARA_TABELA": "eSocial", "REGRA": "eSocial", "SEQUENCIAL": "eSocial",
    "FIXO": "Fixo", "VAZIO": "Fixo",
    "SEM_ORIGEM": "Cliente", "RELATORIO": "Cliente", "PENDENTE": "Cliente",
}
# Quem precisa agir no complementar. 'Legado' entra quando a coluna da
# planilha existir: e o cliente que traz do sistema antigo.
PREENCHE_CLIENTE = {"cliente", "legado"}
# Identidade do colaborador copiada do 1012 nos layouts 100% complementares.
IDENTIDADE = ("NUMEMP", "TIPCOL", "NUMCAD")

AZUL = "1F3864"
AMARELO = "FFE699"
LARANJA = "F8CBAD"
CINZA = "EDF1F5"
FONTE = "Arial"


def papel(campo):
    """Quem preenche o campo: a coluna da planilha, ou o derivado da acao."""
    declarado = (campo.get("quem_preenche") or "").strip()
    if declarado:
        return declarado
    return PAPEL_POR_ACAO.get(campo.get("acao", ""), "Cliente")


def colaboradores(linhas_1012, parametros):
    """[{NUMEMP, TIPCOL, NUMCAD}] distintos, na ordem do 1012."""
    nomes = [c["campo"] for c in parametros["modulos"]["1012"]["campos"]]
    vistos, saida = set(), []
    for linha in linhas_1012 or []:
        ident = tuple(linha[nomes.index(k)] if k in nomes else "" for k in IDENTIDADE)
        if ident[2] and ident not in vistos:
            vistos.add(ident)
            saida.append(dict(zip(IDENTIDADE, ident)))
    return saida


def gerar(resultados, parametros, linhas_1012=None):
    """xlsx com uma aba por layout. resultados: {cod: {"linhas": [...]}}.

    linhas_1012: linhas do cadastro de colaborador, para pre-preencher os
    layouts 100% complementares. Sem elas, usa resultados["1012"] se houver.

    Devolve (bytes, resumo) -- resumo: {cod: (linhas, colunas_do_cliente,
    celulas_obrigatorias_vazias)}.
    """
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # Modo streaming: as linhas vao para o disco conforme sao escritas, em vez
    # de ficarem todas em memoria ate o save. Na cliente piloto sao milhoes de celulas.
    # Os estilos sao criados uma vez so e reaproveitados em todas as celulas.
    lado = Side(style="thin", color="C6CFD8")
    borda = Border(left=lado, right=lado, top=lado, bottom=lado)
    f_amarelo = PatternFill("solid", fgColor=AMARELO)
    f_azul = PatternFill("solid", fgColor=AZUL)
    f_laranja = PatternFill("solid", fgColor=LARANJA)
    f_cinza = PatternFill("solid", fgColor=CINZA)
    fonte_cab_cli = Font(name=FONTE, bold=True, color="000000")
    fonte_cab = Font(name=FONTE, bold=True, color="FFFFFF")
    fonte_quem = Font(name=FONTE, size=9, italic=True)
    fonte_info = Font(name=FONTE, size=9, color="44546A")
    centro = Alignment(horizontal="center")

    wb = openpyxl.Workbook(write_only=True)
    capa = wb.create_sheet("Como preencher")
    capa.column_dimensions["A"].width = 100
    textos = [
        ("Layouts complementares — Migração Senior HCM", True),
        ("", False),
        ("Cada aba é um layout de importação do Senior, já com os dados extraídos do eSocial.", False),
        ("Colunas com cabeçalho AMARELO: o eSocial não traz o dado — preencha.", False),
        ("Células LARANJA: campo obrigatório que deveria vir do eSocial e não veio neste registro — complete.", False),
        ("Não altere as demais colunas nem a ordem das colunas.", False),
        ("A linha 2 de cada aba diz quem preenche cada campo; a linha 3, se é obrigatório e o formato.", False),
    ]
    for t, negrito in textos:
        cel = WriteOnlyCell(capa, value=t)
        cel.font = Font(name=FONTE, size=14 if negrito else 11, bold=negrito)
        capa.append([cel])

    # Atribuir fonte/fundo/borda celula a celula faz o openpyxl calcular o hash
    # dos objetos de estilo toda vez -- era 3/4 do tempo. Cada combinacao e
    # registrada uma vez e as celulas recebem so a copia do indice.
    from copy import copy
    moldes = {}

    def celula(ws, valor, fonte=None, fundo=None, alinhamento=None, formato=None):
        chave = (id(fonte), id(fundo), id(alinhamento), formato)
        molde = moldes.get(chave)
        if molde is None:
            molde = WriteOnlyCell(ws)
            molde.border = borda
            if fonte is not None:
                molde.font = fonte
            if fundo is not None:
                molde.fill = fundo
            if alinhamento is not None:
                molde.alignment = alinhamento
            if formato is not None:
                molde.number_format = formato
            moldes[chave] = molde
        cel = WriteOnlyCell(ws, value=valor)
        cel._style = copy(molde._style)
        return cel

    resumo = {}
    for cod in sorted(resultados):
        mod = parametros["modulos"][cod]
        campos = mod["campos"]
        linhas = resultados[cod].get("linhas") or []
        ws = wb.create_sheet(("%s %s" % (cod, mod.get("nome", "")))[:31])
        papeis = [papel(c) for c in campos]
        nomes = [c["campo"] for c in campos]
        if (not linhas and "NUMCAD" in nomes
                and all(p.lower() in PREENCHE_CLIENTE for p in papeis)):
            if linhas_1012 is None:
                linhas_1012 = (resultados.get("1012") or {}).get("linhas")
            pessoas = colaboradores(linhas_1012, parametros)
            linhas = [[p[n] if n in IDENTIDADE else "" for n in nomes] for p in pessoas]
            if pessoas:
                papeis = ["Cadastro (1012)" if n in IDENTIDADE else p
                          for n, p in zip(nomes, papeis)]
        cliente = [p.lower() in PREENCHE_CLIENTE for p in papeis]

        # write_only: larguras e congelamento antes da primeira linha
        for j, c in enumerate(campos, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(
                12, min(40, len(c.get("descricao") or "") + 2))
        ws.freeze_panes = "A5"

        ws.append([celula(ws, c["campo"], fonte_cab_cli if cliente[j] else fonte_cab,
                          f_amarelo if cliente[j] else f_azul, centro)
                   for j, c in enumerate(campos)])
        ws.append([celula(ws, papeis[j], fonte_quem, f_cinza)
                   for j in range(len(campos))])
        ws.append([celula(ws, "%s · %s" % (
                       "Obrigatório" if c.get("obrigatorio") else "Opcional",
                       c.get("mascara") or c.get("tamanho") or ""), fonte_info, f_cinza)
                   for c in campos])
        cab4 = []
        for c in campos:
            cel = WriteOnlyCell(ws, value=c.get("descricao"))
            cel.font = fonte_info
            cab4.append(cel)
        ws.append(cab4)

        # fundo fixo por coluna; so a celula obrigatoria vazia muda por linha
        falta_possivel = [(not cliente[j]) and bool(c.get("obrigatorio"))
                          and papeis[j] == "eSocial" for j, c in enumerate(campos)]
        faltas = 0
        for linha in linhas:
            saida_linha = []
            for j, valor in enumerate(linha):
                vazio = valor in ("", None)
                fundo = None
                if cliente[j]:
                    fundo = f_amarelo
                elif vazio and falta_possivel[j]:
                    fundo = f_laranja
                    faltas += 1
                saida_linha.append(celula(ws, None if vazio else valor,
                                          fundo=fundo, formato="@"))
            ws.append(saida_linha)
        resumo[cod] = (len(linhas), sum(cliente), faltas)

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue(), resumo


# ---- volta do complementar preenchido ------------------------------------
# O cliente devolve o mesmo xlsx que o extrator gerou, completado. Cada aba vira
# o TXT oficial do layout, sem ler XML de novo: o que veio do eSocial ja esta
# nas linhas. O xlsx e a fonte; aqui so se le, normaliza o que o Excel mudou
# (numero, data) e aponta o que ainda falta. Nada e completado por conta propria.
LINHA_CAMPOS = 1
PRIMEIRA_LINHA_DADOS = 5


def _texto_da_celula(valor, campo):
    """Valor da celula -> texto no formato do layout."""
    import datetime
    import writer
    if valor is None:
        return ""
    if isinstance(valor, (datetime.datetime, datetime.date)):
        # o Excel transforma em data o que o cliente digita como data
        if (campo.get("mascara") or "").upper() == "MM/YYYY":
            return valor.strftime("%m/%Y")
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, bool):
        return "S" if valor else "N"
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)          # 1.0 digitado vira '1', nao '1.0'
    return writer.formatar(str(valor).strip(), campo)


def ler_preenchido(dados_xlsx, parametros):
    """xlsx do complementar preenchido -> {cod: {"linhas", "avisos", "faltas"}}.

    A aba e reconhecida pelo codigo no inicio do nome ('1012 Cadastro...') e as
    colunas pelo nome do campo na linha 1, nao pela posicao. Coluna do layout
    ausente sai vazia e vira aviso; coluna desconhecida e ignorada com aviso.
    faltas: {campo: quantidade de linhas com obrigatorio vazio}.
    """
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(dados_xlsx), read_only=True, data_only=True)
    saida = {}
    for ws in wb.worksheets:
        cod = ws.title[:4]
        if cod not in parametros["modulos"]:
            continue
        campos = parametros["modulos"][cod]["campos"]
        linhas_aba = ws.iter_rows(values_only=True)
        cab = next(linhas_aba, None) or ()
        cab = [str(x).strip().upper() if x is not None else "" for x in cab]
        posicao = {nome: j for j, nome in enumerate(cab) if nome}
        avisos = []
        ausentes = [c["campo"] for c in campos if c["campo"] not in posicao]
        if ausentes:
            avisos.append("Coluna(s) ausente(s), saem vazias: %s" % ", ".join(ausentes))
        sobra = sorted(set(posicao) - {c["campo"] for c in campos})
        if sobra:
            avisos.append("Coluna(s) que o layout nao tem, ignoradas: %s" % ", ".join(sobra))

        linhas, faltas = [], {}
        for n, bruta in enumerate(linhas_aba, start=2):
            if n < PRIMEIRA_LINHA_DADOS:
                continue            # linhas 2-4: quem preenche, formato, descricao
            if not bruta or all(v is None or str(v).strip() == "" for v in bruta):
                continue
            linha = []
            for c in campos:
                j = posicao.get(c["campo"])
                v = _texto_da_celula(bruta[j] if j is not None and j < len(bruta) else None, c)
                if v == "" and c.get("obrigatorio"):
                    faltas[c["campo"]] = faltas.get(c["campo"], 0) + 1
                linha.append(v)
            linhas.append(linha)
        saida[cod] = {"linhas": linhas, "avisos": avisos, "faltas": faltas}
    wb.close()
    return saida
