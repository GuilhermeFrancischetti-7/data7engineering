# -*- coding: utf-8 -*-
"""
Converte a planilha de parametro (Leiautes_HCM_Parametro_Extrator_vN.xlsx) no
arquivo parametros.json, que e o que o extrator le em tempo de execucao.

Rode este script de novo quando sair uma versao nova da planilha (v10, v11...):
    py gerar_parametros.py "caminho\\da\\planilha_v10.xlsx"

O extrator NAO le a planilha direto de proposito: assim a versao usada fica
congelada num arquivo versionavel, e da para ver no diff o que mudou entre v9
e v10 antes de a mudanca entrar em producao.
"""
import json
import os
import re
import sys

import openpyxl

# A planilha mora dentro do projeto, em parametro\. Uma copia so: havendo
# outra em Downloads ou no Drive, a edicao vai para uma e o extrator le a
# outra. Caminho relativo a este arquivo, entao mover o projeto de maquina
# nao quebra nada.
def _planilha_mais_recente():
    """A maior versao presente na pasta parametro.

    Nao fixa o nome: apontar para uma versao antiga fazia o script, rodado sem
    argumento, regredir o mapeamento em silencio. Aconteceu -- uma execucao
    reconstruiu o parametros.json a partir da v11, desfazendo a V12 e a V13.
    """
    pasta = os.path.join(os.path.dirname(os.path.abspath(__file__)), "parametro")
    achados = []
    for nome in os.listdir(pasta) if os.path.isdir(pasta) else []:
        m = re.match(r"^Leiautes_HCM_Parametro_Extrator_v(\d+)\.xlsx$", nome)
        if m and not nome.startswith("~$"):
            achados.append((int(m.group(1)), os.path.join(pasta, nome)))
    return max(achados)[1] if achados else ""


PADRAO_PLANILHA = _planilha_mais_recente()
DESTINO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "parametros.json")

# Colunas da planilha -> chave no JSON.
COLUNAS = {
    "Campo": "campo",
    "Descrição": "descricao",
    "Obrigatório": "obrigatorio",
    "Campo Chave": "chave",
    "Lista": "lista",
    "Tipo": "tipo",
    "Tamanho": "tamanho",
    "Máscara": "mascara",
    "DE/PARA": "depara_schema",
    "Obs. Senior": "obs_senior",
    "Observações (Anterior)": "obs_anterior",
    "▶ Evento": "evento",
    "▶ Caminho no XML": "caminho",
    "▶ Regra / transformação": "regra",
    "▶ Tipo De/Para": "tipo_depara",
    "▶ Autoria da definição": "autoria",
    "◆ % na massa": "pct_massa",
    "◆ Situação": "situacao",
    "◆ Sugestão / origem da decisão": "sugestao",
    # Opcional: eSocial / Cliente / Fixo / Legado / Senior. Ausente, o
    # complementar deriva o papel da acao (complementar.papel).
    "▶ Quem preenche": "quem_preenche",
}

PARADAS = ("menu senior", "linha exemplo", "objetivo", "layout", "tabela",
           "pré-requisito", "pre-requisito", "dica")


def texto(v):
    if v is None:
        return ""
    return str(v).replace("\n", " ").strip()


def classificar(tipo_depara, obs_anterior, regra, caminho):
    """Reduz as varias colunas a UMA acao que o writer sabe executar.

    Precedencia: a coluna '▶ Tipo De/Para' e a definicao aplicada e manda em
    tudo. So quando ela delega ('conforme Anterior') e que a observacao Anterior
    decide. Nada aqui adivinha correspondencia de codigo -- quando falta
    definicao, devolve PENDENTE e o motivo vai para o relatorio.
    """
    t = (tipo_depara or "").strip().lower()
    o = (obs_anterior or "").strip().lower()
    r = (regra or "").strip()

    def valor_fixo():
        m = re.search(r'valor fixo\s*"?([^"]*)"?', r, re.I) or \
            re.search(r'valor fixo\s*"?([^"]*)"?', obs_anterior or "", re.I)
        return (m.group(1).strip() if m else "", bool(m))

    if t.startswith("de-para [cliente]"):
        # 'De-Para [cliente] - como NUMCAD': o campo consulta a tabela de OUTRO
        # campo. O CADATU e o FICREG do 1021 sao o codigo do colaborador, o
        # mesmo do NUMCAD, e a tabela do cliente esta gravada sob NUMCAD.
        m = re.search(r"como\s+([A-Z][A-Z0-9]{2,})", tipo_depara or "", re.I)
        return "DEPARA_CLIENTE", (m.group(1).upper() if m else "")
    if t.startswith("de-para [geral]"):
        return "DEPARA_GERAL", ""
    if t.startswith("de-para [tabela]"):
        m = re.search(r"leiaute\s*(\d+)", tipo_depara, re.I)
        return "DEPARA_TABELA", (m.group(1) if m else "")
    if t == "sem de/para - direto":
        return "DIRETO", ""
    if t == "sem de/para - valor fixo":
        v, ok = valor_fixo()
        return ("FIXO", v) if ok else ("PENDENTE", "Tipo 'valor fixo' sem constante declarada na planilha.")
    if t == "sem de/para - sequencial":
        return "SEQUENCIAL", ""
    if t == "sem de/para - regra":
        return "REGRA", r
    if t == "a definir":
        return "PENDENTE", "Marcado como 'a definir' na planilha de parametro."

    # 'conforme Anterior' / 'sem De/Para (conforme Anterior)' -> a observacao Anterior decide.
    if "conforme anterior" in t or t in ("", "—", "-"):
        # Anterior descreve a numeracao em texto livre, sem marcar o tipo:
        # "Numerar os dependentes..." / "linhas numeradas, sequenciadas...".
        if re.search(r"numerar|numerad|sequenciad", o, re.I):
            return "SEQUENCIAL", (obs_anterior or "").strip()
        if "enviar vazio" in o:
            return "VAZIO", ""
        if "valor fixo" in o:
            v, ok = valor_fixo()
            return ("FIXO", v) if ok else ("PENDENTE", "Anterior diz 'Valor Fixo' sem declarar a constante.")
        if "fonte: relat" in o:
            return "RELATORIO", ""
        if "de-para [cliente]" in o:
            return "DEPARA_CLIENTE", ""
        if "de-para [geral]" in o:
            return "DEPARA_GERAL", ""
        if "de-para [tabela]" in o:
            m = re.search(r"leiaute\s*(\d+)", obs_anterior, re.I)
            return "DEPARA_TABELA", (m.group(1) if m else "")
        if caminho:
            return "DIRETO", ""
        if "enviar vazio" in (r or "").lower():
            return "VAZIO", ""
        return "SEM_ORIGEM", ""

    return "PENDENTE", "Tipo De/Para nao reconhecido: " + texto(tipo_depara)


def main():
    origem = sys.argv[1] if len(sys.argv) > 1 else PADRAO_PLANILHA
    if not os.path.exists(origem):
        print("ERRO: planilha nao encontrada:", origem)
        return 1

    wb = openpyxl.load_workbook(origem, read_only=True, data_only=True)
    modulos = {}
    total = 0

    # Eventos que o extrator le sem preencher coluna de leiaute -- hoje so o
    # S-1010, cuja descricao de rubrica vai para a planilha de De/Para. Como
    # nenhum campo os declara, a raiz deles nao seria deduzida dos caminhos e o
    # leitor descartaria os arquivos na carga.
    auxiliares = {}
    if "Eventos auxiliares" in wb.sheetnames:
        for r in wb["Eventos auxiliares"].iter_rows(values_only=True):
            if not r or len(r) < 2:
                continue
            evento, raiz = texto(r[0]), texto(r[1])
            if re.match(r"^S-\d{4}$", evento) and raiz.startswith("evt"):
                auxiliares[raiz] = evento

    for ws in wb.worksheets:
        titulo = ws.title.strip()
        if not titulo[:4].isdigit():
            continue
        codigo = titulo[:4]
        nome = titulo[7:].strip() if " - " in titulo else titulo

        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(linhas)
                   if r and texto(r[0]) == "Campo"), None)
        if hi is None:
            continue

        cab = [texto(x) for x in linhas[hi]]
        idx = {}
        for rotulo, chave in COLUNAS.items():
            if rotulo in cab:
                idx[chave] = cab.index(rotulo)
        # colunas de evento do Anterior: cabecalho no formato 'S-2200'
        cols_evento = [j for j, h in enumerate(cab) if re.match(r"^S-\d{4}$", h)]

        # tabela do banco Senior, quando declarada no topo da aba
        tabela = ""
        for r in linhas[:hi]:
            if r and texto(r[0]).lower().startswith("tabela"):
                tabela = texto(r[1]) if len(r) > 1 else ""
                break

        # Precedencia entre eventos, declarada no topo da aba no formato
        # 'S-2300 > S-2306 por cpfTrab': quando os dois eventos falam da mesma
        # pessoa, o da esquerda vence e o da direita nao gera linha. Fica na
        # planilha porque e regra de negocio, nao de leitura de XML.
        precedencia = []
        for r in linhas[:hi]:
            if not r or not texto(r[0]).lower().startswith(("precedência de evento",
                                                            "◆ precedência de evento",
                                                            "precedencia de evento",
                                                            "◆ precedencia de evento")):
                continue
            for parte in (texto(r[1]) if len(r) > 1 else "").split(";"):
                m = re.match(r"^\s*(S-\d{4})\s*>\s*(S-\d{4})\s+por\s+(\S+)\s*$", parte)
                if m:
                    precedencia.append({"vence": m.group(1), "perde": m.group(2),
                                        "por": m.group(3)})
            break

        # Blocos que o documento precisa ter para gerar linha, declarados no topo
        # da aba. O 1031 e o caso: 'evtAdmissao/trabalhador/dependente'. Sem
        # isso, todo S-2200 gerava uma linha de dependente, mesmo o do
        # funcionario que nao tem nenhum -- so a chave preenchida.
        #
        # Nao da para deduzir do XML: quem tem UM dependente e quem nao tem
        # NENHUM sao indistinguiveis para a deteccao de repeticao, que so
        # reconhece o bloco quando ele aparece mais de uma vez.
        exige_bloco = []
        for r in linhas[:hi]:
            if r and texto(r[0]).lower().lstrip("◆ ").startswith("exige bloco"):
                exige_bloco = [b.strip()
                               for b in (texto(r[1]) if len(r) > 1 else "").split("|")
                               if b.strip()]
                break

        # Leiaute cujo conteudo NAO sai do XML: e uma tabela fixa do Senior,
        # declarada numa aba 'Lista ...'. O 1006 e o caso -- os motivos de
        # reajuste salarial sao sempre os mesmos e o eSocial nao transmite
        # motivo de alteracao salarial.
        linhas_fixas = ""
        for r in linhas[:hi]:
            if r and texto(r[0]).lower().lstrip("◆ ").startswith("linhas fixas"):
                linhas_fixas = texto(r[1]) if len(r) > 1 else ""
                break

        # Quais eventos GERAM linha. Os demais declarados na aba entram so como
        # fonte de consulta. O 1001 e o caso: a linha e o estabelecimento, que
        # vem do S-1005; o S-1000 so empresta a razao social. Sem isto o S-1000
        # virava uma linha de filial sem filial nenhuma.
        eventos_linha = []
        for r in linhas[:hi]:
            if r and texto(r[0]).lower().lstrip("◆ ").startswith("evento que gera linha"):
                eventos_linha = [e.strip()
                                 for e in (texto(r[1]) if len(r) > 1 else "").split("|")
                                 if e.strip()]
                break

        # '◆ Filtro: caminho = valor | caminho = valor' -- so gera linha o
        # documento cujo caminho tem o valor (1033: codMotAfast = 15).
        filtro = []
        for r in linhas[:hi]:
            if r and texto(r[0]).lower().lstrip("◆ ").startswith("filtro"):
                for parte in (texto(r[1]) if len(r) > 1 else "").split("|"):
                    if "=" in parte:
                        cam, val = parte.split("=", 1)
                        filtro.append([cam.strip(), val.strip()])
                break

        # '◆ Equivalencia de chave: origem -> destino | ...' -- a mesma coisa
        # identificada por duas chaves. O cargo do eSocial simplificado vem so
        # com nmCargo+CBOCargo, e o mesmo cargo pode estar no S-1030 pelo
        # codCargo: sem isto ganhava dois numeros no 1004. O destino e a chave
        # que vale; a origem e lida nos documentos do evento do destino.
        # '◆ Numerar por grupo: CAMPO | ...' -- o sequencial de identidade desses
        # campos recomeca em 1 a cada valor da PRIMEIRA parte da chave composta.
        # O CODBAI do 1011 e numerado por cidade (chave codMunic+bairro), como o
        # consultor pediu, e continua sendo identidade: o 1013 acha o mesmo numero.
        numerar_por_grupo = []
        for r in linhas[:hi]:
            if r and texto(r[0]).lower().lstrip("◆ ").startswith("numerar por grupo"):
                numerar_por_grupo = [c.strip() for c in (texto(r[1]) if len(r) > 1 else "").split("|")
                                     if c.strip()]
                break

        equivalencia = []
        for r in linhas[:hi]:
            if r and texto(r[0]).lower().lstrip("◆ ").startswith(("equivalência de chave",
                                                                  "equivalencia de chave")):
                for parte in (texto(r[1]) if len(r) > 1 else "").split("|"):
                    if "->" in parte:
                        de, para = parte.split("->", 1)
                        equivalencia.append([de.strip(), para.strip()])
                break

        campos = []
        for r in linhas[hi + 1:]:
            if not r or not r[0]:
                continue
            nome_campo = texto(r[0])
            if nome_campo.lower().startswith(PARADAS):
                break

            reg = {}
            for chave, j in idx.items():
                reg[chave] = texto(r[j]) if j < len(r) else ""
            reg["campo"] = nome_campo

            # As colunas por evento do mapeamento Anterior (S-2200, S-2206, ...)
            # trazem a tag do MESMO campo em cada evento. A coluna aplicada
            # guarda so um caminho; guardar as tags permite montar a linha a
            # partir de um evento diferente sem inventar caminho.
            reg["tags_evento"] = {
                cab[j]: texto(r[j])
                for j in cols_evento
                if j < len(r) and texto(r[j])
            }

            # Um campo pode ser lido de mais de um evento. A coluna Evento
            # aceita a MESMA lista separada por '|' da coluna de caminho, na
            # mesma ordem: 'S-2200 | S-2300 | S-2306' pareia com os tres
            # caminhos declarados ao lado. E a unica forma de dizer a que
            # evento pertence uma raiz que so aparece como alternativa --
            # evtTSVAltContr nao tem como ser deduzida de lugar nenhum.
            #
            # 'evento' (singular) continua sendo o primeiro da lista, para nao
            # mexer em nada que ja lia esse campo.
            reg["eventos"] = [e.strip() for e in reg.get("evento", "").split("|")
                              if e.strip()]
            reg["evento"] = reg["eventos"][0] if reg["eventos"] else ""

            acao, arg = classificar(reg.get("tipo_depara"), reg.get("obs_anterior"),
                                    reg.get("regra"), reg.get("caminho"))
            reg["acao"] = acao
            reg["arg"] = arg
            # 'De-Para [tabela] - Leiaute 1012 por cpfTrab': quando o evento nao
            # traz o dado no caminho declarado, o writer busca o MESMO campo no
            # leiaute de origem pela tag de juncao (o S-2299 nao tem codCateg; a
            # admissao da mesma pessoa tem).
            m = re.search(r"\bpor\s+(\w+)\s*$", reg.get("tipo_depara", ""))
            reg["juncao"] = m.group(1) if acao == "DEPARA_TABELA" and m else ""
            # 'De-Para [geral] - por caminho': a tabela De/Para e escolhida pela
            # TAG lida, e nao pelo nome do campo Senior. O eSocial renomeou tags
            # na virada para o S-1.3 e os dois dominios usam os MESMOS numeros
            # com significados diferentes (VISEST: classTrabEstrang 2 = Visto
            # temporario, condIng 2 = Solicitante de refugio). Uma tabela por tag
            # evita que as duas respostas virem uma so.
            reg["por_caminho"] = bool(
                acao in ("DEPARA_GERAL", "DEPARA_CLIENTE")
                and re.search(r"\bpor\s+caminho\s*$",
                              reg.get("tipo_depara", ""), re.I))
            reg["obrigatorio"] = reg.get("obrigatorio", "").lower().startswith("sim")
            reg["chave"] = reg.get("chave", "").lower().startswith("sim")
            campos.append(reg)
            total += 1

        if campos:
            modulos[codigo] = {"codigo": codigo, "nome": nome, "tabela": tabela,
                               "precedencia": precedencia,
                               "exige_bloco": exige_bloco,
                               "eventos_linha": eventos_linha,
                               "filtro": filtro,
                               "equivalencia": equivalencia,
                               "numerar_por_grupo": numerar_por_grupo,
                               "linhas_fixas": linhas_fixas, "campos": campos}

    # ---- mascara a mostrar na planilha de De/Para ----
    # A planilha pergunta uma vez por campo, valendo para todos os leiautes.
    # Sete campos tem mascara diferente conforme o leiaute, e em dois deles o
    # TAMANHO difere (CODCID 5 ou 7, CODSIN 2 ou 4): preenchido pelo menor, o
    # outro leiaute recusaria. Esta aba diz qual mostrar. A coluna Mascara de
    # cada leiaute fica intacta -- e a que o Senior publica.
    mascaras_depara = {}
    if "Mascara De-Para" in wb.sheetnames:
        for r in wb["Mascara De-Para"].iter_rows(values_only=True):
            if not r or len(r) < 2:
                continue
            campo, mascara = texto(r[0]), texto(r[1])
            if campo.isupper() and len(campo) >= 4 and mascara:
                mascaras_depara[campo] = mascara

    # ---- de qual tabela do eSocial vem a descricao de cada campo ----
    # Serve para a planilha de De/Para mostrar "15 - Gozo de ferias" em vez de
    # so "15". Nao traduz nada; e texto de apoio para quem preenche.
    tabelas_por_campo = {}
    if "Tabelas eSocial por campo" in wb.sheetnames:
        for r in wb["Tabelas eSocial por campo"].iter_rows(values_only=True):
            if not r or len(r) < 2:
                continue
            campo, numero = texto(r[0]), texto(r[1])
            if campo.isupper() and len(campo) >= 4 and numero.isdigit():
                tabelas_por_campo[campo] = numero.zfill(2)

    # ---- listas de dominio SO do lado Senior ----
    # Nao sao De/Para: o lado eSocial e codigo do cliente e so ele sabe parear.
    # Servem para a planilha de De/Para oferecer as opcoes validas do Senior em
    # vez de uma celula em branco. Hoje: os 382 eventos padrao do CODEVE.
    listas_senior = {}
    for ws in wb.worksheets:
        if not ws.title.startswith("Lista "):
            continue
        alvo = ws.title.replace("Lista ", "").split()[0].strip()
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(linhas)
                   if r and any(texto(x) == "Código" for x in r if x)), None)
        if hi is None:
            continue
        cab = [texto(x) for x in linhas[hi]]
        cc = cab.index("Código")
        itens = []
        for r in linhas[hi + 1:]:
            if not r or cc >= len(r) or not texto(r[cc]):
                continue
            itens.append({
                "codigo": texto(r[cc]),
                "descricao": texto(r[cc + 1]) if cc + 1 < len(r) else "",
                "detalhe": texto(r[cc + 2]) if cc + 2 < len(r) else "",
                "tipo": texto(r[cc + 3]) if cc + 3 < len(r) else "",
            })
        if itens:
            listas_senior[alvo] = itens

    # ---- tabelas De-Para de dominio que ja vem prontas na planilha ----
    deparas = {}
    for ws in wb.worksheets:
        if not ws.title.startswith("De-Para "):
            continue
        alvo = ws.title.replace("De-Para ", "").strip()
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(linhas)
                   if r and any(texto(x) == "Código eSocial" for x in r if x)), None)
        if hi is None:
            continue
        cab = [texto(x) for x in linhas[hi]]
        ce = cab.index("Código eSocial")
        cs = cab.index("Código Senior") if "Código Senior" in cab else ce + 2
        de = cab.index("Descrição eSocial") if "Descrição eSocial" in cab else ce + 1
        ds = cab.index("Descrição Senior") if "Descrição Senior" in cab else cs + 1
        itens = []
        for r in linhas[hi + 1:]:
            if not r or ce >= len(r) or not r[ce]:
                continue
            itens.append({
                "esocial": texto(r[ce]),
                "desc_esocial": texto(r[de]) if de < len(r) else "",
                "senior": texto(r[cs]) if cs < len(r) else "",
                "desc_senior": texto(r[ds]) if ds < len(r) else "",
            })
        # A mesma tabela pode valer para mais de um campo com o mesmo dominio:
        # 'De-Para CODPAI PAINAS' (pais de nacionalidade e de nascimento).
        if itens:
            for campo_alvo in alvo.split():
                deparas[campo_alvo] = itens

    for cod, mod in modulos.items():
        aba = mod.get("linhas_fixas")
        if not aba:
            continue
        if aba not in wb.sheetnames:
            print("AVISO: %s declara linhas fixas em '%s', que nao existe." % (cod, aba))
            mod["linhas_fixas"] = []
            continue
        brutas = [list(r) for r in wb[aba].iter_rows(values_only=True)]
        hi2 = next((i for i, r in enumerate(brutas)
                    if r and any(texto(x) == "Código" for x in r if x)), None)
        corpo = []
        for r in brutas[(hi2 or 0) + 1:]:
            valores = [texto(x) for x in r][:len(mod["campos"])]
            if not valores or not valores[0]:
                continue
            valores += [""] * (len(mod["campos"]) - len(valores))
            corpo.append(valores)
        mod["linhas_fixas"] = corpo
        print("Linhas fixas de %s: %d, de '%s'" % (cod, len(corpo), aba))

    wb.close()

    saida = {
        "origem": os.path.basename(origem),
        "versao": re.search(r"v(\d+)", os.path.basename(origem), re.I).group(1)
                  if re.search(r"v(\d+)", os.path.basename(origem), re.I) else "?",
        "modulos": modulos,
        "deparas_dominio": deparas,
        "eventos_auxiliares": auxiliares,
        "listas_senior": listas_senior,
        "tabelas_por_campo": tabelas_por_campo,
        "mascaras_depara": mascaras_depara,
    }
    with open(DESTINO, "w", encoding="utf-8") as f:
        json.dump(saida, f, ensure_ascii=False, indent=1)

    print("origem :", os.path.basename(origem))
    print("destino:", DESTINO)
    print("modulos:", len(modulos), "| campos:", total)
    print("De-Para de dominio prontos:", {k: len(v) for k, v in deparas.items()})
    print("Listas so do lado Senior:", {k: len(v) for k, v in listas_senior.items()})
    print("Campos com tabela do eSocial:", tabelas_por_campo)
    print("Mascaras fixadas para o De/Para:", mascaras_depara)
    from collections import Counter
    c = Counter(cp["acao"] for m in modulos.values() for cp in m["campos"])
    print("acoes:", dict(c.most_common()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
